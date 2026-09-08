import json
from pathlib import Path
from pptx import Presentation

REF=Path('2. Metrics Review Apr - Jun 26 NEW FORMAT.pptx')
OUT=Path('analysis/reference_structure.json')
OUT.parent.mkdir(parents=True,exist_ok=True)

def shape_info(sh):
    d={'name':sh.name,'type':str(sh.shape_type),'left':sh.left,'top':sh.top,'width':sh.width,'height':sh.height}
    if getattr(sh,'has_text_frame',False):
        d['text']=sh.text
    if getattr(sh,'has_table',False):
        d['table']=[[c.text for c in row.cells] for row in sh.table.rows]
    if getattr(sh,'has_chart',False):
        ch=sh.chart
        cd={'chart_type':str(ch.chart_type),'has_title':ch.has_title}
        if ch.has_title: cd['title']=ch.chart_title.text_frame.text
        series=[]
        for s in ch.series:
            sd={'name':s.name}
            try: sd['values']=list(s.values)
            except: pass
            try: sd['categories']=[str(c) for c in s.categories]
            except: pass
            series.append(sd)
        cd['series']=series
        d['chart']=cd
    if str(sh.shape_type).startswith('GROUP'):
        d['children']=[shape_info(c) for c in sh.shapes]
    return d

prs=Presentation(REF)
data={'slide_width':prs.slide_width,'slide_height':prs.slide_height,'slides':[]}
for i,slide in enumerate(prs.slides,1):
    data['slides'].append({'slide':i,'shapes':[shape_info(s) for s in slide.shapes]})
OUT.write_text(json.dumps(data,indent=2,ensure_ascii=False,default=str),encoding='utf-8')
print(f'wrote {OUT}')
