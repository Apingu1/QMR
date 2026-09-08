import json, os, re, zipfile
from datetime import datetime, timedelta
from pathlib import Path
import xml.etree.ElementTree as ET
from pptx import Presentation

NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships','pr':'http://schemas.openxmlformats.org/package/2006/relationships'}
DATE_BUILTINS=set(range(14,23))|{45,46,47}

def col_num(ref):
    m=re.match(r'([A-Z]+)', ref or '')
    if not m:return 0
    n=0
    for c in m.group(1): n=n*26+ord(c)-64
    return n

def is_date_fmt(code):
    if not code:return False
    c=re.sub(r'"[^"]*"','',code.lower()); c=re.sub(r'\\.','',c)
    return bool(re.search(r'(^|[^a-z])[ymdhis]+', c))

def excel_date(v):
    try:
        d=datetime(1899,12,30)+timedelta(days=float(v))
        return d.strftime('%Y-%m-%d') if d.time()==datetime.min.time() else d.strftime('%Y-%m-%d %H:%M:%S')
    except: return v

def parse_xlsx(path):
    with zipfile.ZipFile(path) as z:
        shared=[]
        if 'xl/sharedStrings.xml' in z.namelist():
            root=ET.fromstring(z.read('xl/sharedStrings.xml'))
            for si in root.findall('m:si',NS): shared.append(''.join(t.text or '' for t in si.iterfind('.//m:t',NS)))
        date_styles=set()
        if 'xl/styles.xml' in z.namelist():
            root=ET.fromstring(z.read('xl/styles.xml'))
            custom={int(x.attrib['numFmtId']):x.attrib.get('formatCode','') for x in root.findall('.//m:numFmts/m:numFmt',NS)}
            cellxfs=root.find('m:cellXfs',NS)
            if cellxfs is not None:
                for i,xf in enumerate(cellxfs.findall('m:xf',NS)):
                    nid=int(xf.attrib.get('numFmtId','0'))
                    if nid in DATE_BUILTINS or is_date_fmt(custom.get(nid,'')): date_styles.add(i)
        wb=ET.fromstring(z.read('xl/workbook.xml'))
        relroot=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        rels={x.attrib['Id']:x.attrib['Target'] for x in relroot.findall('pr:Relationship',NS)}
        out={'file':Path(path).name,'sheets':[]}
        for s in wb.findall('.//m:sheets/m:sheet',NS):
            name=s.attrib['name']; rid=s.attrib['{'+NS['r']+'}id']; target=rels[rid]
            sp='xl/'+target.lstrip('/') if not target.startswith('xl/') else target
            root=ET.fromstring(z.read(sp)); cells={}; maxc=maxr=0
            for c in root.findall('.//m:sheetData/m:row/m:c',NS):
                ref=c.attrib.get('r',''); rm=re.search(r'(\d+)$',ref)
                if not rm: continue
                row=int(rm.group(1)); col=col_num(ref); maxc=max(maxc,col); maxr=max(maxr,row)
                t=c.attrib.get('t'); style=int(c.attrib.get('s','0')); v=c.find('m:v',NS); inline=c.find('m:is',NS); f=c.find('m:f',NS); val=''
                if t=='s' and v is not None:
                    try: val=shared[int(v.text)]
                    except: val=v.text or ''
                elif t=='inlineStr' and inline is not None: val=''.join(x.text or '' for x in inline.iterfind('.//m:t',NS))
                elif t=='b' and v is not None: val=(v.text=='1')
                elif v is not None:
                    raw=v.text or ''
                    if style in date_styles: val=excel_date(raw)
                    else:
                        try:
                            fv=float(raw); val=int(fv) if fv.is_integer() else fv
                        except: val=raw
                cells[(row,col)]={'v':val,'f':f.text if f is not None else None}
            rows=[]
            for r in range(1,maxr+1):
                vals=[cells.get((r,c),{}).get('v','') for c in range(1,maxc+1)]
                if any(x not in ('',None) for x in vals): rows.append({'row':r,'values':vals})
            out['sheets'].append({'name':name,'max_row':maxr,'max_col':maxc,'rows':rows})
        return out

def dump_pptx(path):
    p=Presentation(path); slides=[]
    for i,sl in enumerate(p.slides,1):
        shapes=[]
        for j,sh in enumerate(sl.shapes,1):
            text=sh.text if hasattr(sh,'text') else ''
            shapes.append({'index':j,'name':sh.name,'type':str(sh.shape_type),'left':sh.left,'top':sh.top,'width':sh.width,'height':sh.height,'text':text})
        slides.append({'slide':i,'shapes':shapes})
    return {'file':Path(path).name,'width':p.slide_width,'height':p.slide_height,'slides':slides}

outdir=Path('analysis/extracted'); outdir.mkdir(parents=True,exist_ok=True)
files=sorted(Path('0. QMS Spreadsheets 2026').glob('*.xlsx'))+[Path('Root Cause Analysis Data & Trending (Version 2).xlsx')]
for p in files:
    (outdir/(p.stem+'.json')).write_text(json.dumps(parse_xlsx(p),indent=2,ensure_ascii=False),encoding='utf-8')
ref=Path('2. Metrics Review Apr - Jun 26 NEW FORMAT.pptx')
(outdir/'reference_pptx_layout.json').write_text(json.dumps(dump_pptx(ref),indent=2,ensure_ascii=False),encoding='utf-8')
print('extracted',len(files),'workbooks')
