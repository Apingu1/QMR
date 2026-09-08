import json, re
from pathlib import Path
from datetime import datetime, date
from collections import Counter

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.dml.color import RGBColor

ROOT = Path('.')
DATA = json.loads((ROOT / 'analysis/qmr_analysis.json').read_text(encoding='utf-8'))
OUTDIR = ROOT / 'output'
OUTDIR.mkdir(exist_ok=True)
OUT = OUTDIR / 'Eaststone_QMR_Apr-Jun_2026.pptx'

BG = RGBColor(248, 248, 246)
NAVY = RGBColor(21, 41, 64)
TEAL = RGBColor(34, 111, 121)
GREEN = RGBColor(36, 130, 70)
AMBER = RGBColor(210, 137, 32)
RED = RGBColor(165, 50, 50)
GREY = RGBColor(90, 95, 101)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(30, 30, 30)

SYSTEM_ORDER = [
    'Deviations', 'GMP Complaints', 'Supplier Complaints', 'OOS-OOT',
    'CAPA', 'Change Controls', 'FMEA', 'Incidents',
    'Document Change Requests', 'Effectiveness Checks', 'MDR'
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def sys(name):
    return DATA['systems'].get(name, {'Q1': {'counts': {}, 'records': [], 'avg_closure_days': 0}, 'Q2': {'counts': {}, 'records': [], 'avg_closure_days': 0, 'significant': []}})


def clean(v):
    return ' '.join(str(v or '').replace('\n', ' ').split())


def fmt_date(v):
    if not v:
        return ''
    try:
        return datetime.strptime(str(v)[:10], '%Y-%m-%d').strftime('%d-%b-%y')
    except Exception:
        return str(v)


def short(v, n=120):
    s = clean(v)
    return s if len(s) <= n else s[: n - 1] + '…'


def add_bg(slide, section=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shp.fill.solid(); shp.fill.fore_color.rgb = BG
    shp.line.fill.background()
    if section:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.72), prs.slide_height)
        bar.fill.solid(); bar.fill.fore_color.rgb = TEAL
        bar.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(0.04), Inches(5.95), Inches(0.65), Inches(1.35))
        tf = tb.text_frame; tf.clear()
        p = tf.paragraphs[0]; p.text = section.upper(); p.alignment = PP_ALIGN.CENTER
        r = p.runs[0]; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE


def add_title(slide, title, subtitle=None):
    tb = slide.shapes.add_textbox(Inches(0.95), Inches(0.22), Inches(11.9), Inches(0.6))
    tf = tb.text_frame; tf.clear()
    p = tf.paragraphs[0]; p.text = title
    r = p.runs[0]; r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = NAVY
    if subtitle:
        tb2 = slide.shapes.add_textbox(Inches(0.96), Inches(0.82), Inches(11.8), Inches(0.35))
        tf2 = tb2.text_frame; tf2.clear(); p2 = tf2.paragraphs[0]; p2.text = subtitle
        rr = p2.runs[0]; rr.font.size = Pt(14); rr.font.color.rgb = GREY


def add_box(slide, x, y, w, h, title, body, color=NAVY):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = RGBColor(220, 224, 226); shp.line.width = Pt(0.8)
    tb = slide.shapes.add_textbox(Inches(x+0.18), Inches(y+0.12), Inches(w-0.36), Inches(h-0.22))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title
    r = p.runs[0]; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = color
    p2 = tf.add_paragraph(); p2.text = str(body)
    r2 = p2.runs[0]; r2.font.size = Pt(23); r2.font.bold = True; r2.font.color.rgb = BLACK
    return shp


def add_table(slide, x, y, w, h, rows, font_size=10):
    if not rows:
        rows = [['No data']]
    max_cols = max(len(r) for r in rows)
    rows = [list(r) + [''] * (max_cols - len(r)) for r in rows]
    tbl_shape = slide.shapes.add_table(len(rows), max_cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = tbl_shape.table
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = str(value if value is not None else '')
            cell.margin_left = Inches(0.04); cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            fill = cell.fill; fill.solid()
            fill.fore_color.rgb = NAVY if i == 0 else (RGBColor(232, 238, 241) if i % 2 == 0 else WHITE)
            for para in cell.text_frame.paragraphs:
                para.alignment = PP_ALIGN.LEFT
                for run in para.runs:
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = WHITE if i == 0 else BLACK
                    run.font.bold = bool(i == 0)
    return tbl_shape


def add_chart(slide, x, y, w, h, title, cats, series, chart_type=XL_CHART_TYPE.LINE_MARKERS):
    cd = CategoryChartData(); cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    chart = slide.shapes.add_chart(chart_type, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    chart.has_title = True; chart.chart_title.text_frame.text = title
    chart.has_legend = True; chart.legend.position = XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout = False
    return chart


def assessment(q1, q2):
    a = q1['counts'].get('raised', 0) or 0
    b = q2['counts'].get('raised', 0) or 0
    s1 = q1['counts'].get('major_critical_raised', 0) or 0
    s2 = q2['counts'].get('major_critical_raised', 0) or 0
    if b < a and s2 <= s1:
        return 'Improving'
    if b > a * 1.35 or s2 > s1 + 1:
        return 'Deteriorating'
    return 'Stable'


def traffic(value):
    return GREEN if value == 'Improving' else (RED if value == 'Deteriorating' else AMBER)


def month_trend(name):
    q = sys(name)
    seen = {}
    for r in q['Q1'].get('records', []) + q['Q2'].get('records', []):
        seen.setdefault(r.get('ref',''), r)
    cats = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun']
    new = [0] * 6; closed = [0] * 6
    for r in seen.values():
        try:
            d = datetime.strptime(r.get('opened','')[:10], '%Y-%m-%d').date()
            if d.year == 2026 and 1 <= d.month <= 6:
                new[d.month-1] += 1
        except Exception:
            pass
        try:
            d = datetime.strptime(r.get('closed','')[:10], '%Y-%m-%d').date()
            if d.year == 2026 and 1 <= d.month <= 6:
                closed[d.month-1] += 1
        except Exception:
            pass
    return cats, new, closed


def severity_count(q, quarter_start, quarter_end):
    c = Counter()
    for r in q.get('records', []):
        try:
            d = datetime.strptime(r.get('opened','')[:10], '%Y-%m-%d').date()
        except Exception:
            continue
        if quarter_start <= d <= quarter_end:
            sev = clean(r.get('classification')) or 'Other'
            c[sev] += 1
    return c


def sig_rows(q2, max_rows=7):
    sig = q2.get('significant', []) or []
    rows = [['Reference', 'Date', 'Classification', 'Status', 'Summary / Description']]
    if not sig:
        rows.append(['', '', '', '', 'No Major or Critical items identified for Q2.'])
        return rows
    for r in sig[:max_rows]:
        rows.append([r.get('ref',''), fmt_date(r.get('date','')), r.get('reclassification') or r.get('classification',''), r.get('status_end',''), short(r.get('summary',''), 95)])
    if len(sig) > max_rows:
        rows.append(['', '', '', '', f'+{len(sig)-max_rows} additional significant item(s) listed in source data.'])
    return rows


def trend_rows(name, q2, max_rows=7):
    area = Counter(); root = Counter()
    for r in q2.get('records', []):
        if r.get('area'):
            area[short(r.get('area'), 45)] += 1
        roots = re.split(r'[\n;/]+', clean(r.get('root_cause','')))
        for x in roots:
            x = clean(x)
            if x:
                root[short(x, 50)] += 1
    rows = [['Theme', 'Count', 'Comment']]
    for k, v in (root.most_common(6) or area.most_common(6)):
        rows.append([k, v, 'Monitored through routine QMS trending'])
    if len(rows) == 1:
        rows.append(['No recurring trend identified', 0, 'Continue routine monitoring'])
    return rows

# Title
slide = prs.slides.add_slide(blank); add_bg(slide)
add_title(slide, 'QMR for Period April 26 - Jun 26', 'Quality Metrics Review')
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(1.55), Inches(11.45), Inches(0.05)).fill.solid()
slide.shapes[-1].fill.fore_color.rgb = TEAL
for i, (t, b, c) in enumerate([
    ('Review basis', 'QMS source spreadsheets from repository', NAVY),
    ('Period', 'April 2026 - June 2026', TEAL),
    ('Prepared as', 'Management review PowerPoint', AMBER),
]):
    add_box(slide, 1.1 + i*4.0, 2.3, 3.55, 1.45, t, b, c)

# Key
slide = prs.slides.add_slide(blank); add_bg(slide, 'QMR')
add_title(slide, 'QMR Key')
add_table(slide, 1.05, 1.25, 11.6, 3.1, [
    ['Indicator', 'Meaning'],
    ['Green / Improving', 'Metric movement is favourable or quality signal has reduced.'],
    ['Amber / Stable', 'Metric is broadly stable or requires routine monitoring.'],
    ['Red / Deteriorating', 'Metric has increased materially or significant items have increased.'],
    ['Significant items', 'Major and Critical records from the source QMS logs.'],
    ['Record counting', 'Parent QMS records are counted; action rows such as /A, /B are not double counted.'],
], 13)

# Executive summary
slide = prs.slides.add_slide(blank); add_bg(slide, 'Summary')
add_title(slide, 'Executive Summary (QMS Health Overview)')
summary = []
for name in SYSTEM_ORDER:
    q = sys(name); q1, q2 = q['Q1'], q['Q2']
    summary.append([name, q2['counts'].get('raised',0), q2['counts'].get('closed',0), q2['counts'].get('open_end',0), q2['counts'].get('major_critical_raised',0), assessment(q1,q2)])
add_table(slide, 1.0, 1.1, 11.9, 4.4, [['System', 'Raised Q2', 'Closed Q2', 'Open at 30-Jun', 'Major/Critical raised', 'Assessment']] + summary, 10)
add_box(slide, 1.0, 5.85, 11.9, 0.9, 'Overall management conclusion', 'QMS activity remains reviewable through the supplied logs. Key open/significant records remain tracked within the relevant QMS systems.', TEAL)

# Individual systems
for name in SYSTEM_ORDER:
    q = sys(name); q1, q2 = q['Q1'], q['Q2']
    slide = prs.slides.add_slide(blank); add_bg(slide, name)
    add_title(slide, f'{name} - KPI Summary', 'Q1 2026 vs Q2 2026')
    ass = assessment(q1, q2)
    add_box(slide, 1.0, 1.05, 2.2, 1.0, 'Assessment', ass, traffic(ass))
    add_box(slide, 3.45, 1.05, 2.1, 1.0, 'Raised Q2', q2['counts'].get('raised', 0), TEAL)
    add_box(slide, 5.8, 1.05, 2.1, 1.0, 'Closed Q2', q2['counts'].get('closed', 0), GREEN)
    add_box(slide, 8.15, 1.05, 2.1, 1.0, 'Open at 30 Jun', q2['counts'].get('open_end', 0), AMBER)
    add_box(slide, 10.5, 1.05, 2.1, 1.0, 'Major/Critical', q2['counts'].get('major_critical_raised', 0), RED if q2['counts'].get('major_critical_raised',0) else GREEN)
    kpi_rows = [
        ['KPI', 'Q1 2026', 'Q2 2026'],
        ['Total raised this quarter', q1['counts'].get('raised',0), q2['counts'].get('raised',0)],
        ['Total closed this quarter', q1['counts'].get('closed',0), q2['counts'].get('closed',0)],
        ['Major/Critical raised', q1['counts'].get('major_critical_raised',0), q2['counts'].get('major_critical_raised',0)],
        ['Major/Critical closed', q1['counts'].get('major_critical_closed',0), q2['counts'].get('major_critical_closed',0)],
        ['Average closure time', f"{q1.get('avg_closure_days') or 0:g} days", f"{q2.get('avg_closure_days') or 0:g} days"],
        ['Overdue open at quarter end', q1['counts'].get('overdue_open',0), q2['counts'].get('overdue_open',0)],
    ]
    add_table(slide, 1.0, 2.35, 4.2, 3.85, kpi_rows, 10)
    cats, new, closed = month_trend(name)
    add_chart(slide, 5.55, 2.32, 7.1, 3.9, f'{name} - 2026 Monthly Trend', cats, [('New', new), ('Closed', closed)], XL_CHART_TYPE.LINE_MARKERS)

    slide = prs.slides.add_slide(blank); add_bg(slide, name)
    add_title(slide, f'{name} - Significant Items')
    add_table(slide, 1.0, 1.1, 11.9, 4.2, sig_rows(q2, 8), 9)
    add_box(slide, 1.0, 5.65, 11.9, 0.95, 'Quarter position', f"Open at 30 Jun: {q2['counts'].get('open_end',0)} | Overdue open: {q2['counts'].get('overdue_open',0)} | Closed during Q2: {q2['counts'].get('closed',0)}", TEAL)

    slide = prs.slides.add_slide(blank); add_bg(slide, name)
    add_title(slide, f'{name} - Trend Analysis')
    add_table(slide, 1.0, 1.15, 11.9, 4.35, trend_rows(name, q2), 10)
    add_box(slide, 1.0, 5.75, 11.9, 0.85, 'Management review note', 'No automatic escalation is proposed from trend count alone; significant records remain subject to normal QMS ownership and due-date management.', AMBER)

# Closing slide
slide = prs.slides.add_slide(blank); add_bg(slide)
add_title(slide, 'Any Questions?')
add_box(slide, 1.0, 2.3, 11.7, 1.2, 'QMR for Period April 26 - Jun 26', 'Prepared from the repository QMS source spreadsheets.', TEAL)

prs.save(OUT)
print(f'Wrote {OUT}')
